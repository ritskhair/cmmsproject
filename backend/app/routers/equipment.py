from datetime import date, datetime
import io
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_account, require_roles
from app.models.component import Component
from app.models.equipment import Equipment
from app.models.section import Section as SectionModel
from app.schemas.equipment import (
    ComponentCreate,
    ComponentResponse,
    ComponentUpdate,
    EquipmentCreate,
    EquipmentImportSummary,
    EquipmentResponse,
    EquipmentStatus,
    EquipmentUpdate,
    SectionResponse,
)

router = APIRouter(prefix="/equipment", tags=["equipment"])

MANAGEMENT_ROLES = ("teknisi", "team_leader", "manager", "general_manager", "super_admin")
HIERARCHY_READ_ROLES = (*MANAGEMENT_ROLES, "operator")


@router.get("/sections", response_model=list[SectionResponse])
def list_sections(
    db: Session = Depends(get_db),
    _=Depends(require_roles(*HIERARCHY_READ_ROLES)),
):
    return db.query(SectionModel).filter(SectionModel.active.is_(True)).order_by(SectionModel.code).all()


@router.get("/sections/{section_id}/machines", response_model=list[EquipmentResponse])
def list_section_machines(
    section_id: UUID,
    db: Session = Depends(get_db),
    _=Depends(require_roles(*HIERARCHY_READ_ROLES)),
):
    return db.query(Equipment).filter(Equipment.section_id == section_id).order_by(Equipment.asset_id).all()


@router.get("/machines/{machine_id}/components", response_model=list[ComponentResponse])
def list_machine_components(
    machine_id: UUID,
    db: Session = Depends(get_db),
    _=Depends(require_roles(*HIERARCHY_READ_ROLES)),
):
    if not db.query(Equipment).filter(Equipment.id == machine_id).first():
        raise HTTPException(status_code=404, detail="Machine not found")
    return db.query(Component).filter(Component.machine_id == machine_id).order_by(Component.created_at.desc()).all()


@router.post("/machines/{machine_id}/components", response_model=ComponentResponse, status_code=201)
def create_component(
    machine_id: UUID,
    payload: ComponentCreate,
    db: Session = Depends(get_db),
    _=Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if not db.query(Equipment).filter(Equipment.id == machine_id).first():
        raise HTTPException(status_code=404, detail="Machine not found")
    component = Component(machine_id=machine_id, **payload.model_dump())
    db.add(component)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Failed to create component")
    db.refresh(component)
    return component


@router.patch("/components/{component_id}", response_model=ComponentResponse)
def update_component(
    component_id: UUID,
    payload: ComponentUpdate,
    db: Session = Depends(get_db),
    _=Depends(require_roles(*MANAGEMENT_ROLES)),
):
    component = db.query(Component).filter(Component.id == component_id).first()
    if not component:
        raise HTTPException(status_code=404, detail="Component not found")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(component, key, value)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Failed to update component")
    db.refresh(component)
    return component


@router.delete("/components/{component_id}", status_code=204)
def delete_component(
    component_id: UUID,
    db: Session = Depends(get_db),
    _=Depends(require_roles("super_admin")),
):
    component = db.query(Component).filter(Component.id == component_id).first()
    if not component:
        raise HTTPException(status_code=404, detail="Component not found")
    db.delete(component)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Component is referenced by a work order")


@router.get("", response_model=list[EquipmentResponse])
def list_equipment(
    status: EquipmentStatus | None = Query(default=None),
    section: str | None = Query(default=None),
    db: Session = Depends(get_db),
    _=Depends(require_roles("teknisi", "team_leader", "manager", "general_manager", "super_admin")),
):
    query = db.query(Equipment)
    if status:
        query = query.filter(Equipment.status == status)
    if section:
        query = query.filter(Equipment.section == section)
    return query.order_by(Equipment.asset_id).all()


@router.post("", response_model=EquipmentResponse, status_code=201)
def create_equipment(payload: EquipmentCreate, db: Session = Depends(get_db), _=Depends(require_roles("teknisi", "team_leader", "manager", "general_manager", "super_admin"))):
    values = payload.model_dump()
    section = db.query(SectionModel).filter(SectionModel.code == payload.section.value).first()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")
    equipment = Equipment(**values, section_id=section.id)
    db.add(equipment)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Asset ID already exists")
    db.refresh(equipment)
    return equipment


def normalize_section_value(val: any) -> str:
    if val is None:
        return "1"
    s = str(val).strip().lower()
    for d in ["1", "2", "3", "4", "5"]:
        if d in s:
            return d
    return "1"


def parse_date_value(val: any) -> datetime | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day)
    try:
        val_str = str(val).strip()
        if not val_str:
            return None
        if len(val_str) == 10:
            return datetime.strptime(val_str, "%Y-%m-%d")
        return datetime.fromisoformat(val_str)
    except Exception:
        return None


def generate_equipment_template_bytes() -> bytes:
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    guide_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    italic_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Sheet Guide
    ws_guide = wb.active
    ws_guide.title = "Panduan & Ketentuan"
    ws_guide.views.sheetView[0].showGridLines = True
    ws_guide["A1"] = "PANDUAN PENGISIAN IMPORT EQUIPMENT & KOMPONEN"
    ws_guide["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    ws_guide["A2"] = "(Sheet ini TIDAK PERLU DIHAPUS saat upload, sistem otomatis membaca sheet Equipment & Components)"
    ws_guide["A2"].font = italic_font

    instructions = [
        ("No", "Aturan / Fitur", "Penjelasan Lengkap"),
        ("1", "Nama Sheet & Header", "JANGAN mengubah nama sheet 'Equipment' dan 'Components', serta JANGAN mengubah teks baris header (Baris 1)."),
        ("2", "Sheet 'Equipment'", "Wajib diisi: 'Asset ID', 'Nama Mesin', 'Tipe Mesin', 'Section' (Section 1 s/d Section 5).\nStatus bersifat opsional (jika kosong otomatis 'Operational')."),
        ("3", "Sheet 'Components'", "Kolom 'Asset ID Mesin' wajib diisi sama persis dengan Asset ID di sheet Equipment agar komponen langsung terhubung ke mesinnya."),
        ("4", "Tipe / Merek Komponen", "Bebas diisi merek atau tipe komponen (misal: Siemens, Mitsubishi, Omron, SMC, FESTO, dsb)."),
        ("5", "Kolom Last PM", "Opsional. Jika dikosongkan tidak apa-apa, nanti akan otomatis terupdate saat ada riwayat EWO/PM baru."),
        ("6", "Metode Upload di Web", "Bisa Drag & Drop file langsung ke browser atau klik tombol Browse/Pilih File."),
    ]

    for row_idx, row_data in enumerate(instructions, start=4):
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws_guide.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if row_idx == 4:
                cell.fill = guide_header_fill
                cell.font = header_font
                cell.alignment = center_align
            else:
                cell.font = bold_font if col_idx in (1, 2) else regular_font
                cell.alignment = center_align if col_idx == 1 else left_align

    ws_guide.column_dimensions["A"].width = 8
    ws_guide.column_dimensions["B"].width = 25
    ws_guide.column_dimensions["C"].width = 88

    # Sheet Equipment
    ws_eq = wb.create_sheet(title="Equipment")
    ws_eq.views.sheetView[0].showGridLines = True
    eq_headers = ["Asset ID *", "Nama Mesin *", "Tipe Mesin *", "Section *", "Status (Opsional)", "Last PM (Opsional)"]
    for col_idx, header in enumerate(eq_headers, start=1):
        cell = ws_eq.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    sample_eq_data = [
        ("EQ-001", "CNC Milling Machine 01", "Machining", "Section 1", "Operational", "2026-09-15"),
        ("EQ-002", "Hydraulic Press 50T", "Pressing", "Section 2", "", ""),
        ("EQ-003", "Main Conveyor Line A", "Conveyor", "Section 3", "Operational", ""),
        ("EQ-004", "Air Compressor Unit 2", "Utility", "Section 1", "", ""),
        ("EQ-005", "Automated Stamping Press", "Stamping", "Section 4", "", ""),
    ]
    for row_idx, row_data in enumerate(sample_eq_data, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws_eq.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.alignment = center_align if col_idx in (1, 4, 5, 6) else left_align
            cell.border = thin_border

    for col in ws_eq.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_eq.column_dimensions[col_letter].width = max(max_len + 4, 18)

    # Sheet Components
    ws_comp = wb.create_sheet(title="Components")
    ws_comp.views.sheetView[0].showGridLines = True
    comp_headers = ["Asset ID Mesin *", "Nama Komponen *", "Tipe / Merek Komponen", "Status Komponen (Opsional)"]
    for col_idx, header in enumerate(comp_headers, start=1):
        cell = ws_comp.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    sample_comp_data = [
        ("EQ-001", "Spindle Motor 7.5kW", "Mitsubishi", "Operational"),
        ("EQ-001", "Coolant Circulation Pump", "Grundfos", ""),
        ("EQ-001", "Ball Screw X-Axis", "THK", ""),
        ("EQ-002", "Main Hydraulic Cylinder", "Rexroth", ""),
        ("EQ-002", "Hydraulic Power Pack Unit", "Yuken", ""),
        ("EQ-002", "Safety Light Curtain Sensor", "Keyence", ""),
        ("EQ-003", "Conveyor Drive Motor 3HP", "Toshiba", ""),
        ("EQ-003", "Rubber Belt Heavy Duty", "Habasit", ""),
        ("EQ-004", "Screw Air End Unit", "Atlas Copco", ""),
        ("EQ-004", "Intake Air Filter Housing", "SMC", ""),
        ("EQ-005", "Pneumatic Clutch & Brake", "FESTO", ""),
        ("EQ-005", "PLC Control Module FX-3U", "Mitsubishi", ""),
    ]
    for row_idx, row_data in enumerate(sample_comp_data, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws_comp.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.alignment = center_align if col_idx in (1, 3, 4) else left_align
            cell.border = thin_border

    for col in ws_comp.columns:
        max_len = max(len(str(cell.value or "")) for col_cell in col for cell in [col_cell])
        col_letter = get_column_letter(col[0].column)
        ws_comp.column_dimensions[col_letter].width = max(max_len + 4, 24)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


@router.get("/import-template")
def get_import_template(
    _=Depends(require_roles("super_admin")),
):
    excel_bytes = generate_equipment_template_bytes()
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_import_equipment.xlsx"},
    )


@router.post("/import-excel", response_model=EquipmentImportSummary)
async def import_equipment_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _=Depends(require_roles("super_admin")),
):
    filename = file.filename or ""
    if not (filename.endswith(".xlsx") or filename.endswith(".xlsm")):
        raise HTTPException(status_code=400, detail="File harus berformat Excel (.xlsx)")

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file Excel: {str(e)}")

    # Find sheets
    sheet_names_lower = {name.strip().lower(): name for name in wb.sheetnames}
    
    eq_sheet_name = None
    for name_lower, actual_name in sheet_names_lower.items():
        if "equipment" in name_lower or "mesin" in name_lower:
            eq_sheet_name = actual_name
            break
            
    comp_sheet_name = None
    for name_lower, actual_name in sheet_names_lower.items():
        if "component" in name_lower or "komponen" in name_lower:
            comp_sheet_name = actual_name
            break

    # Fallback to sheet indices if not found by name
    if not eq_sheet_name:
        for actual_name in wb.sheetnames:
            if "panduan" not in actual_name.lower() and "ketentuan" not in actual_name.lower():
                eq_sheet_name = actual_name
                break

    if not eq_sheet_name:
        raise HTTPException(status_code=400, detail="Sheet data 'Equipment' tidak ditemukan di dalam file Excel.")

    sections = db.query(SectionModel).all()
    section_map_by_code = {s.code: s for s in sections}
    section_map_by_name = {s.name.lower(): s for s in sections}

    total_equipment_created = 0
    total_equipment_updated = 0
    total_components_added = 0
    errors: list[str] = []

    # Map asset_id (clean lowercase) -> Equipment instance
    equipment_by_asset_id: dict[str, Equipment] = {}

    # 1. Process Equipment Sheet
    ws_eq = wb[eq_sheet_name]
    eq_header_map: dict[str, int] = {}
    
    # Read headers from row 1
    for col_idx in range(1, ws_eq.max_column + 1):
        cell_val = ws_eq.cell(row=1, column=col_idx).value
        if cell_val:
            cleaned = str(cell_val).strip().lower().replace("*", "").strip()
            eq_header_map[cleaned] = col_idx

    # Find column indices
    def find_eq_col(*candidates: str) -> int | None:
        for c in candidates:
            for header, idx in eq_header_map.items():
                if c in header:
                    return idx
        return None

    asset_col = find_eq_col("asset id", "asset_id", "kode")
    name_col = find_eq_col("nama mesin", "machine name", "name", "nama")
    type_col = find_eq_col("tipe mesin", "machine type", "type", "tipe")
    section_col = find_eq_col("section", "area", "seksi")
    status_col = find_eq_col("status")
    last_pm_col = find_eq_col("last pm", "last_pm", "pm terakhir")

    if not asset_col or not name_col:
        raise HTTPException(
            status_code=400,
            detail="Header kolom wajib ('Asset ID' dan 'Nama Mesin') tidak ditemukan pada sheet Equipment.",
        )

    for row_idx in range(2, ws_eq.max_row + 1):
        asset_id_val = ws_eq.cell(row=row_idx, column=asset_col).value if asset_col else None
        if not asset_id_val or not str(asset_id_val).strip():
            continue

        asset_id = str(asset_id_val).strip()
        name_val = ws_eq.cell(row=row_idx, column=name_col).value if name_col else None
        name = str(name_val).strip() if name_val else f"Machine {asset_id}"

        type_val = ws_eq.cell(row=row_idx, column=type_col).value if type_col else None
        type_str = str(type_val).strip() if type_val else "General"

        section_raw = ws_eq.cell(row=row_idx, column=section_col).value if section_col else "1"
        section_code = normalize_section_value(section_raw)

        # Match section_id
        matching_section = section_map_by_code.get(section_code)
        if not matching_section:
            matching_section = section_map_by_name.get(f"section {section_code}")

        status_val = ws_eq.cell(row=row_idx, column=status_col).value if status_col else None
        status_str = str(status_val).strip() if status_val else "Operational"
        if status_str not in ("Operational", "Warning", "Critical", "Under Maintenance", "Breakdown"):
            status_str = "Operational"

        last_pm_val = ws_eq.cell(row=row_idx, column=last_pm_col).value if last_pm_col else None
        last_pm_dt = parse_date_value(last_pm_val)

        # Check if equipment exists
        existing_eq = db.query(Equipment).filter(Equipment.asset_id == asset_id).first()
        if existing_eq:
            existing_eq.name = name
            existing_eq.type = type_str
            existing_eq.section = section_code
            if matching_section:
                existing_eq.section_id = matching_section.id
            if status_str:
                existing_eq.status = status_str
            if last_pm_dt is not None:
                existing_eq.last_pm = last_pm_dt
            equipment_by_asset_id[asset_id.lower()] = existing_eq
            total_equipment_updated += 1
        else:
            new_eq = Equipment(
                asset_id=asset_id,
                name=name,
                type=type_str,
                section=section_code,
                section_id=matching_section.id if matching_section else None,
                status=status_str,
                last_pm=last_pm_dt,
            )
            db.add(new_eq)
            db.flush()  # to generate id
            equipment_by_asset_id[asset_id.lower()] = new_eq
            total_equipment_created += 1

    # 2. Process Components Sheet (if present)
    if comp_sheet_name:
        ws_comp = wb[comp_sheet_name]
        comp_header_map: dict[str, int] = {}
        for col_idx in range(1, ws_comp.max_column + 1):
            cell_val = ws_comp.cell(row=1, column=col_idx).value
            if cell_val:
                cleaned = str(cell_val).strip().lower().replace("*", "").strip()
                comp_header_map[cleaned] = col_idx

        def find_comp_col(*candidates: str) -> int | None:
            for c in candidates:
                for header, idx in comp_header_map.items():
                    if c in header:
                        return idx
            return None

        comp_asset_col = find_comp_col("asset id", "asset_id", "mesin", "machine")
        comp_name_col = find_comp_col("nama komponen", "component name", "komponen", "name")
        comp_type_col = find_comp_col("tipe", "merek", "brand", "type")
        comp_status_col = find_comp_col("status")

        if comp_asset_col and comp_name_col:
            for row_idx in range(2, ws_comp.max_row + 1):
                c_asset_val = ws_comp.cell(row=row_idx, column=comp_asset_col).value
                c_name_val = ws_comp.cell(row=row_idx, column=comp_name_col).value
                if not c_asset_val or not c_name_val:
                    continue

                c_asset_str = str(c_asset_val).strip()
                c_name_str = str(c_name_val).strip()
                c_type_val = ws_comp.cell(row=row_idx, column=comp_type_col).value if comp_type_col else None
                c_type_str = str(c_type_val).strip() if c_type_val else "General"
                c_status_val = ws_comp.cell(row=row_idx, column=comp_status_col).value if comp_status_col else None
                c_status_str = str(c_status_val).strip() if c_status_val else "Operational"
                if c_status_str not in ("Operational", "Warning", "Critical", "Under Maintenance", "Breakdown"):
                    c_status_str = "Operational"

                target_eq = equipment_by_asset_id.get(c_asset_str.lower())
                if not target_eq:
                    target_eq = db.query(Equipment).filter(Equipment.asset_id == c_asset_str).first()
                    if target_eq:
                        equipment_by_asset_id[c_asset_str.lower()] = target_eq

                if not target_eq:
                    errors.append(f"Komponen baris {row_idx} ('{c_name_str}'): Asset ID Mesin '{c_asset_str}' tidak ditemukan.")
                    continue

                # Check if component name already exists for this machine
                existing_comp = (
                    db.query(Component)
                    .filter(Component.machine_id == target_eq.id, Component.name.ilike(c_name_str))
                    .first()
                )
                if existing_comp:
                    existing_comp.type = c_type_str
                    if c_status_str:
                        existing_comp.status = c_status_str
                else:
                    new_comp = Component(
                        name=c_name_str,
                        type=c_type_str,
                        status=c_status_str,
                        machine_id=target_eq.id,
                    )
                    db.add(new_comp)
                    total_components_added += 1

    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Gagal menyimpan data import ke database: {str(e)}")

    return EquipmentImportSummary(
        success=True,
        total_equipment_created=total_equipment_created,
        total_equipment_updated=total_equipment_updated,
        total_components_added=total_components_added,
        errors=errors,
    )


@router.get("/operator/equipment", response_model=list[EquipmentResponse])
def operator_equipment_options(
    section: str = Query(..., pattern="^[1-5]$"),
    db: Session = Depends(get_db),
    _=Depends(require_roles("operator")),
):
    return db.query(Equipment).filter(Equipment.section == section).order_by(Equipment.asset_id).all()



@router.get("/{equipment_id}", response_model=EquipmentResponse)
def get_equipment(equipment_id: UUID, db: Session = Depends(get_db), _=Depends(require_roles("teknisi", "team_leader", "manager", "general_manager", "super_admin"))):
    equipment = db.query(Equipment).filter(Equipment.id == equipment_id).first()
    if not equipment:
        raise HTTPException(status_code=404, detail="Equipment not found")
    return equipment


@router.patch("/{equipment_id}", response_model=EquipmentResponse)
def update_equipment(equipment_id: UUID, payload: EquipmentUpdate, db: Session = Depends(get_db), _=Depends(require_roles("teknisi", "team_leader", "manager", "general_manager", "super_admin"))):
    equipment = db.query(Equipment).filter(Equipment.id == equipment_id).first()
    if not equipment:
        raise HTTPException(status_code=404, detail="Equipment not found")
    for key, value in payload.model_dump(exclude_unset=True).items():
        if key == "section":
            section = db.query(SectionModel).filter(SectionModel.code == value.value).first()
            if not section:
                raise HTTPException(status_code=404, detail="Section not found")
            equipment.section_id = section.id
        setattr(equipment, key, value)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Asset ID already exists")
    db.refresh(equipment)
    return equipment


@router.delete("/{equipment_id}", status_code=204)
def delete_equipment(
    equipment_id: UUID,
    db: Session = Depends(get_db),
    _=Depends(require_roles("super_admin")),
):
    equipment = db.query(Equipment).filter(Equipment.id == equipment_id).first()
    if not equipment:
        raise HTTPException(status_code=404, detail="Equipment not found")
    db.delete(equipment)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Equipment is referenced by an active work order and cannot be deleted")
